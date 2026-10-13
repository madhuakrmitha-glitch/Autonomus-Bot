"""Word report helper (FR10; python-docx, reference S22).

``build_docx`` renders claims with evidence citations; ``validate_docx`` reopens the saved
bytes and checks that every claim is present, every citation marker resolves to a listed
source, every listed source is a real evidence record, and the checksum matches. A narrow,
registered helper: it never receives shell access.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from .contracts import Claim, EvidenceRecord
from .util import display_time, sha256_bytes, truncate

LABEL = re.compile(r"\[(E\d+)\]")


@dataclass
class ReportSection:
    heading: str
    claims: list[Claim] = field(default_factory=list)
    figures: list[tuple[bytes, str, list[str]]] = field(default_factory=list)  # (png bytes, caption, evidence ids)


@dataclass
class ReportInput:
    title: str
    task_id: str
    goal: str
    summary_claims: list[Claim]
    sections: list[ReportSection]
    evidence: list[EvidenceRecord]
    limitations: list[str]
    disagreements: list[str]
    method: dict[str, Any]
    timezone: str
    generated_at: str


def label_map(inp: ReportInput) -> dict[str, str]:
    order: list[str] = []
    for c in [*inp.summary_claims, *(c for s in inp.sections for c in s.claims)]:
        for e in c.evidence_ids:
            if e not in order:
                order.append(e)
    for s in inp.sections:
        for _, _, eids in s.figures:
            for e in eids:
                if e not in order:
                    order.append(e)
    return {eid: f"E{i + 1}" for i, eid in enumerate(order)}


def _claim_paragraph(doc: Any, claim: Claim, labels: dict[str, str], style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.add_run(claim.text.strip())
    cites = "".join(f"[{labels[e]}]" for e in claim.evidence_ids if e in labels)
    if claim.support == "supported" and cites:
        p.add_run(" " + cites)
    else:
        marker = {"unsupported": "unsupported - no cited evidence", "disputed": "disputed",
                  "uncertain": "uncertain"}.get(claim.support, claim.support)
        if cites:
            p.add_run(" " + cites)
        r = p.add_run(f" ({marker}{': ' + claim.review_note if claim.review_note else ''})")
        r.italic = True
        r.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)


def build_docx(inp: ReportInput) -> tuple[bytes, dict[str, str]]:
    labels = label_map(inp)
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    doc.add_heading(inp.title, level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Task {inp.task_id} | generated {display_time(inp.generated_at, inp.timezone)} | "
                 f"{len(labels)} cited source records").italic = True
    doc.add_paragraph(f"Request: {truncate(inp.goal, 600)}")

    doc.add_heading("Summary", level=1)
    if inp.summary_claims:
        for c in inp.summary_claims:
            _claim_paragraph(doc, c, labels)
    else:
        doc.add_paragraph("No supported summary could be produced from the available evidence.")

    for section in inp.sections:
        doc.add_heading(section.heading, level=1)
        for c in section.claims:
            _claim_paragraph(doc, c, labels, style="List Bullet")
        for png, caption, eids in section.figures:
            doc.add_picture(io.BytesIO(png), width=Inches(5.5))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.add_run(caption + " " + "".join(f"[{labels[e]}]" for e in eids if e in labels)).italic = True

    if inp.disagreements:
        doc.add_heading("Where sources disagree", level=1)
        for d in inp.disagreements:
            doc.add_paragraph(d, style="List Bullet")

    doc.add_heading("Limitations and evidence gaps", level=1)
    for lim in inp.limitations or ["None recorded."]:
        doc.add_paragraph(lim, style="List Bullet")

    doc.add_heading("Sources", level=1)
    by_id = {e.evidence_id: e for e in inp.evidence}
    for eid, label in labels.items():
        e = by_id.get(eid)
        if e is None:
            continue
        p = doc.add_paragraph()
        p.add_run(f"[{label}] ").bold = True
        p.add_run(e.source_url)
        details = [f"retrieved {display_time(e.captured_at, inp.timezone)}",
                   f"published {e.source_date}" if e.source_date else "publication date not stated"]
        if e.kind == "video_interval" and e.video_interval:
            details.append(f"video {e.video_interval.start_seconds:.1f}s-{e.video_interval.end_seconds:.1f}s")
        if e.kind == "image_region":
            details.append("image region" + (f" {int(e.region.width)}x{int(e.region.height)}px" if e.region else ""))
        details.append(f"evidence {e.evidence_id}")
        p.add_run(" - " + "; ".join(details))
        if e.excerpt:
            q = doc.add_paragraph()
            q.paragraph_format.left_indent = Inches(0.4)
            q.add_run(f"“{truncate(e.excerpt, 300)}”").italic = True
        if e.limitations:
            lp = doc.add_paragraph()
            lp.paragraph_format.left_indent = Inches(0.4)
            lp.add_run(f"Limitation: {e.limitations}").italic = True

    doc.add_heading("Method", level=1)
    for k, v in inp.method.items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")
    props = doc.core_properties
    props.title = inp.title[:200]
    props.subject = f"OMCA task {inp.task_id}"
    props.keywords = "omca;report"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), labels


HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
RULE = re.compile(r"^([-*_])\1{2,}$")
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
INLINE = re.compile(r"(\*\*[^*\n]+\*\*|`[^`\n]+`|\*[^*\s][^*\n]*\*)")


def _runs(p: Any, text: str) -> None:
    """Add ``text`` to a paragraph, rendering **bold**, *italic* and `code`."""
    for part in INLINE.split(text):
        if not part:
            continue
        if len(part) > 4 and part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif len(part) > 2 and part.startswith("`") and part.endswith("`"):
            p.add_run(part[1:-1]).font.name = "Consolas"
        elif len(part) > 2 and part.startswith("*") and part.endswith("*"):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build_document(title: str, markdown: str, *, task_id: str, generated_at: str, timezone: str) -> bytes:
    """A free-form Word document from simple Markdown: headings, lists, tables, quotes, code, **bold**, *italic*."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=0)
    doc.add_paragraph().add_run(f"Task {task_id} | generated {display_time(generated_at, timezone)}").italic = True
    lines = (markdown or "").replace("\r\n", "\n").split("\n")
    para: list[str] = []
    started = False

    def flush() -> None:
        if para:
            _runs(doc.add_paragraph(), " ".join(para))
            para.clear()

    i = 0
    while i < len(lines):
        line, stripped = lines[i], lines[i].strip()
        if stripped.startswith("```"):
            flush()
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            run = doc.add_paragraph().add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            i += 1
            started = True
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]):
            flush()
            header, rows = _cells(line), []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            for r, cells in enumerate([header, *rows]):
                for col in range(len(header)):
                    p = table.cell(r, col).paragraphs[0]
                    _runs(p, cells[col] if col < len(cells) else "")
                    if r == 0:
                        for run in p.runs:
                            run.bold = True
            started = True
            continue
        heading, bullet, numbered = HEADING.match(stripped), BULLET.match(line), NUMBERED.match(line)
        if heading:
            flush()
            # models often repeat the title as the first heading
            if started or heading.group(2).strip() != title.strip():
                doc.add_heading(heading.group(2), level=min(len(heading.group(1)), 3))
        elif RULE.match(stripped):
            flush()
        elif bullet or numbered:
            flush()
            m = bullet or numbered
            style = ("List Bullet" if bullet else "List Number") + (" 2" if len(m.group(1)) >= 2 else "")
            _runs(doc.add_paragraph(style=style), m.group(2))
        elif stripped.startswith(">"):
            flush()
            _runs(doc.add_paragraph(style="Quote"), stripped.lstrip("> "))
        elif not stripped:
            flush()
        else:
            para.append(stripped)
        started = started or bool(stripped)
        i += 1
    flush()
    doc.core_properties.title = title[:200]
    doc.core_properties.subject = f"OMCA task {task_id}"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def validate_docx(data: bytes, *, expected_sha256: str | None, claims: list[Claim], labels: dict[str, str],
                  evidence_ids: set[str]) -> dict[str, Any]:
    problems: list[str] = []
    try:
        text = docx_text(data)
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "problems": [f"report does not reopen: {exc}"]}
    sha = sha256_bytes(data)
    if expected_sha256 and sha != expected_sha256:
        problems.append("checksum mismatch after save")
    body, _, sources = text.partition("\nSources\n")
    for c in claims:
        if c.text.strip() not in text:
            problems.append(f"claim missing from report: {truncate(c.text, 80)}")
        for e in c.evidence_ids:
            if e in labels and f"[{labels[e]}]" not in body:
                problems.append(f"citation {labels[e]} missing for claim {truncate(c.text, 60)}")
    used = set(LABEL.findall(body))
    listed = set(re.findall(r"^\[(E\d+)\] ", sources, re.M))
    for lab in used - listed:
        problems.append(f"citation [{lab}] has no source entry")
    reverse = {v: k for k, v in labels.items()}
    for lab in listed:
        eid = reverse.get(lab)
        if eid is None or eid not in evidence_ids:
            problems.append(f"source entry [{lab}] is not a recorded evidence item")
    return {"ok": not problems, "problems": problems, "sha256": sha, "bytes": len(data), "citations_used": len(used),
            "sources_listed": len(listed), "claims_checked": len(claims)}
