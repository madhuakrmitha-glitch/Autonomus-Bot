"""Bounded office adapters (P3; FR24).

Each adapter does one narrow, reviewed thing:

* ``spreadsheet.read`` / ``spreadsheet.write`` work on declared input workbooks with openpyxl
  and always write a *new* copy; replacing an original is a separate approved file action.
* ``email.draft`` / ``calendar.draft`` produce local previews (.eml / .ics) only.
* ``email.send`` / ``calendar.invite`` deliver through the worker's configured SMTP account and
  only run after the coordinator consumed an owner approval bound to the exact draft hash
  and recipients (the policy layer also rejects recipients outside the contact list).
* ``desktop.raw_input`` exists for completeness but is classified UNKNOWN_MUTATION and is
  executed only on Windows after approval.
"""

from __future__ import annotations

import asyncio
import io
import smtplib
import ssl
import sys
from datetime import datetime
from email.message import EmailMessage
from email.parser import BytesParser
from email.policy import default as default_policy
from typing import Any

from ..config import SmtpConfig
from ..contracts import ErrorCode, WorkerJob
from ..secrets import resolve
from ..util import new_id, sha256_bytes
from .agent import HandlerResult, JobContext, JobFailure

MAX_CELLS = 5000


async def _input(ctx: JobContext, artifact_id: str) -> bytes:
    if artifact_id not in ctx.job.declared_inputs:
        raise JobFailure(ErrorCode.POLICY_DENIED, f"{artifact_id} is not a declared input of this job")
    return await ctx.transport.download_input(ctx.job.job_id, artifact_id)


def _sheet(wb, name: str | None):
    if name is None:
        return wb.active
    if name not in wb.sheetnames:
        raise JobFailure(ErrorCode.VERIFICATION_FAILED, f"sheet {name!r} does not exist")
    return wb[name]


def make_office_handlers(smtp: SmtpConfig | None = None) -> dict[str, Any]:
    async def sheet_read(job: WorkerJob, ctx: JobContext) -> HandlerResult:
        from openpyxl import load_workbook

        data = await _input(ctx, job.arguments["artifact_id"])
        try:
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise JobFailure(ErrorCode.UNSUPPORTED_MEDIA, f"not a readable .xlsx workbook: {exc}") from exc
        ws = _sheet(wb, job.arguments.get("sheet"))
        rng = job.arguments.get("cell_range")
        rows = ws[rng] if rng else ws.iter_rows(max_row=200, max_col=50)
        if rng and not isinstance(rows, tuple):
            rows = ((rows,),)
        elif rng and rows and not isinstance(rows[0], tuple):
            rows = (rows,)
        values: dict[str, Any] = {}
        for row in rows:
            for cell in row:
                if len(values) >= MAX_CELLS:
                    break
                if cell.value is not None and hasattr(cell, "coordinate"):
                    v = cell.value
                    values[cell.coordinate] = v.isoformat() if isinstance(v, datetime) else v
        return HandlerResult(result={"sheet": ws.title, "sheets": wb.sheetnames, "values": values,
                                     "input_sha256": sha256_bytes(data)})

    async def sheet_write(job: WorkerJob, ctx: JobContext) -> HandlerResult:
        from openpyxl import load_workbook

        a = job.arguments
        data = await _input(ctx, a["artifact_id"])
        try:
            wb = load_workbook(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise JobFailure(ErrorCode.UNSUPPORTED_MEDIA, f"not a writable .xlsx workbook: {exc}") from exc
        ws = _sheet(wb, a.get("sheet"))
        for coord, value in a["cells"].items():
            ws[coord] = value
        buf = io.BytesIO()
        wb.save(buf)
        out = buf.getvalue()
        name = a.get("output_name") or "updated.xlsx"
        if not name.lower().endswith(".xlsx"):
            name += ".xlsx"
        art = await ctx.upload("workbook", name, out, meta={"parent_artifact_id": a["artifact_id"],
                                                           "cells_written": len(a["cells"]), "sha256": sha256_bytes(out)})
        # verify by reopening what was produced
        check = load_workbook(io.BytesIO(out))
        cws = _sheet(check, a.get("sheet"))
        mismatched = [k for k, v in a["cells"].items() if cws[k].value != v]
        return HandlerResult(result={"artifact_id": art, "cells_written": len(a["cells"]), "sha256": sha256_bytes(out)},
                             artifact_ids=[art], verification={"checked": True, "passed": not mismatched,
                                                               "checks": {"cells": not mismatched}})

    async def email_draft(job: WorkerJob, ctx: JobContext) -> HandlerResult:
        a = job.arguments
        msg = EmailMessage()
        msg["To"] = ", ".join(a["to"])
        msg["From"] = smtp.from_address if smtp else "owner@localhost"
        msg["Subject"] = a["subject"]
        msg["Message-ID"] = f"<{new_id('eml')}@omca>"
        msg["X-OMCA-Draft"] = "local preview - not sent"
        msg.set_content(a["body"])
        data = bytes(msg)
        art = await ctx.upload("draft", "draft.eml", data, media_type="message/rfc822", meta={"sha256": sha256_bytes(data)})
        return HandlerResult(result={"draft_artifact_id": art, "sha256": sha256_bytes(data), "to": a["to"]},
                             artifact_ids=[art])

    async def calendar_draft(job: WorkerJob, ctx: JobContext) -> HandlerResult:
        a = job.arguments
        uid = new_id("ics")

        def ics_time(v: str) -> str:
            dt = datetime.fromisoformat(v.replace("Z", "+00:00"))
            return dt.strftime("%Y%m%dT%H%M%S") + ("Z" if dt.utcoffset() is not None and dt.utcoffset().total_seconds() == 0 else "")

        lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//OMCA//draft//EN", "METHOD:REQUEST", "BEGIN:VEVENT",
                 f"UID:{uid}@omca", f"DTSTAMP:{datetime.utcnow().strftime('%Y%m%dT%H%M%SZ')}",
                 f"DTSTART:{ics_time(a['start'])}", f"DTEND:{ics_time(a['end'])}", f"SUMMARY:{a['title']}"]
        if a.get("location"):
            lines.append(f"LOCATION:{a['location']}")
        if a.get("description"):
            lines.append("DESCRIPTION:" + a["description"].replace("\n", "\\n"))
        lines += [f"ATTENDEE;RSVP=TRUE:mailto:{x}" for x in a.get("attendees", [])]
        lines += ["END:VEVENT", "END:VCALENDAR"]
        data = ("\r\n".join(lines) + "\r\n").encode("utf-8")
        art = await ctx.upload("draft", "invite.ics", data, media_type="text/calendar", meta={"sha256": sha256_bytes(data)})
        return HandlerResult(result={"draft_artifact_id": art, "sha256": sha256_bytes(data)}, artifact_ids=[art])

    def _send(msg: EmailMessage) -> dict[str, Any]:
        if smtp is None:
            raise JobFailure(ErrorCode.UNSUPPORTED_CAPABILITY, "this worker has no SMTP account configured")
        sending = False
        try:
            with smtplib.SMTP(smtp.host, smtp.port, timeout=30) as s:
                if smtp.starttls:
                    s.starttls(context=ssl.create_default_context())
                user, pw = resolve(smtp.username_ref), resolve(smtp.password_ref)
                if user and pw:
                    s.login(user, pw)
                sending = True
                refused = s.send_message(msg)
        except (smtplib.SMTPRecipientsRefused, smtplib.SMTPSenderRefused) as exc:
            raise JobFailure(ErrorCode.SITE_UNAVAILABLE, f"SMTP refused the message (not sent): {exc}") from exc
        except OSError as exc:  # includes smtplib.SMTPException
            if sending:
                # the server may have accepted the message before the error: never resend automatically
                raise JobFailure(ErrorCode.UNCERTAIN_SIDE_EFFECT, f"SMTP error while sending: {exc}",
                                 uncertain=True) from exc
            raise JobFailure(ErrorCode.SITE_UNAVAILABLE, f"SMTP unavailable before sending: {exc}") from exc
        return {"refused": {k: str(v) for k, v in (refused or {}).items()}}

    async def email_send(job: WorkerJob, ctx: JobContext) -> HandlerResult:
        a = job.arguments
        data = await _input(ctx, a["draft_artifact_id"])
        msg = BytesParser(policy=default_policy).parsebytes(data)
        to = [x.strip() for x in str(msg["To"] or "").split(",") if x.strip()]
        if sorted(t.lower() for t in to) != sorted(t.lower() for t in a["to"]):
            raise JobFailure(ErrorCode.POLICY_DENIED, "draft recipients differ from the approved recipients")
        del msg["X-OMCA-Draft"]
        res = await asyncio.to_thread(_send, msg)
        return HandlerResult(result={"sent": True, "to": to, "message_id": str(msg["Message-ID"]), **res,
                                     "draft_sha256": sha256_bytes(data)})

    async def calendar_invite(job: WorkerJob, ctx: JobContext) -> HandlerResult:
        a = job.arguments
        ics = await _input(ctx, a["draft_artifact_id"])
        msg = EmailMessage()
        msg["To"] = ", ".join(a["attendees"])
        msg["From"] = smtp.from_address if smtp else "owner@localhost"
        msg["Subject"] = "Invitation"
        msg.set_content("You are invited. The calendar invitation is attached.")
        msg.add_attachment(ics, maintype="text", subtype="calendar", filename="invite.ics", params={"method": "REQUEST"})
        res = await asyncio.to_thread(_send, msg)
        return HandlerResult(result={"sent": True, "attendees": a["attendees"], **res})

    async def raw_input(job: WorkerJob, ctx: JobContext) -> HandlerResult:
        if sys.platform != "win32":
            raise JobFailure(ErrorCode.UNSUPPORTED_CAPABILITY, "raw desktop input is only available on Windows workers")
        try:
            from pywinauto import keyboard, mouse  # type: ignore[import-not-found]
        except ImportError as exc:
            raise JobFailure(ErrorCode.UNSUPPORTED_CAPABILITY, "pywinauto is not installed (pip install omca[desktop])") from exc
        a = job.arguments
        if a["action"] == "click":
            mouse.click(coords=(int(a["x"]), int(a["y"])))
        elif a["action"] == "type":
            keyboard.send_keys(a.get("text") or "", with_spaces=True, pause=0.02)
        elif a["action"] == "hotkey":
            keyboard.send_keys(a.get("keys") or "")
        elif a["action"] == "scroll":
            mouse.scroll(coords=(int(a.get("x") or 0), int(a.get("y") or 0)), wheel_dist=-3)
        return HandlerResult(result={"performed": a["action"]})

    return {
        "spreadsheet.read": sheet_read,
        "spreadsheet.write": sheet_write,
        "email.draft": email_draft,
        "calendar.draft": calendar_draft,
        "email.send": email_send,
        "calendar.invite": calendar_invite,
        "desktop.raw_input": raw_input,
    }
