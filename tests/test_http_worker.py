"""A worker driven over the real HTTP API (the transport remote P3 workers use)."""

from __future__ import annotations

import io
import os

import httpx
import pytest
from openpyxl import Workbook, load_workbook

from omca.api import create_app
from omca.config import config_from_dict
from omca.contracts import ReceiptStatus, TaskRequest
from omca.coordinator import Coordinator
from omca.util import new_id, now_iso
from omca.worker.agent import HttpTransport, WorkerAgent
from omca.worker.office import make_office_handlers
from omca.workers import job_deadline

from .conftest import base_config


@pytest.mark.acceptance("AT25")
async def test_remote_worker_over_http(tmp_path):
    os.environ["OMCA_T_W1"] = "w1-secret"
    cfg = config_from_dict(base_config(tmp_path, None, workers=[{"id": "vm1", "credential_ref": "env:OMCA_T_W1",
                                                                 "capabilities": ["spreadsheet"]}],
                                       browser={"attachments": []}))
    coord = Coordinator(cfg)
    app = create_app(coord)
    transport = HttpTransport("http://coordinator", "vm1", "w1-secret")
    transport.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://coordinator",
                                         headers={"Authorization": "Bearer w1-secret"})
    agent = WorkerAgent("vm1", transport, workspace=tmp_path / "w", capabilities=["spreadsheet"],
                        handlers=make_office_handlers(None))
    await agent.start()
    await agent.heartbeat()
    assert coord.workers.workers()[0]["status"] == "ONLINE"

    task = coord.tasks.create(TaskRequest(goal="x"), owner=coord.owner_key, workflow="desktop_workflow", limits={})
    wb = Workbook()
    wb.active["A1"] = 1
    buf = io.BytesIO()
    wb.save(buf)
    art = coord.artifacts.put(task.id, "input", "in.xlsx", buf.getvalue()).manifest.artifact_id
    lease = coord.workers.acquire_lease(task.id, "vm1", f"vm1:task:{task.id}")
    from omca.contracts import WorkerJob

    job = WorkerJob(job_id=new_id("job"), task_id=task.id, action_id=new_id("act"), worker_id="vm1", lease_id=lease.lease_id,
                    resource=lease.resource, fencing_token=lease.fencing_token, operation="spreadsheet.write",
                    arguments={"artifact_id": art, "cells": {"B1": 2}, "output_name": "out.xlsx"},
                    idempotency_key=new_id("k"), deadline=job_deadline(30), lease_expires_at=lease.expires_at,
                    issued_at=now_iso(), declared_inputs=[art])
    coord.workers.create_job(job)
    receipt = await agent.run_once(wait_seconds=1)
    assert receipt is not None and receipt.status == ReceiptStatus.SUCCEEDED, receipt
    stored = coord.workers.receipt_for_action(job.action_id)
    assert stored is not None and stored.status == ReceiptStatus.SUCCEEDED
    out = coord.artifacts.get(stored.result["artifact_id"])
    assert out.sha256 == stored.result["sha256"]
    assert load_workbook(io.BytesIO(coord.artifacts.read_bytes(out.artifact_id))).active["B1"].value == 2
    # a wrong credential is refused before anything is claimed
    bad = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://coordinator",
                            headers={"Authorization": "Bearer nope"})
    r = await bad.post("/workers/vm1/jobs", json={"session_epoch": agent.epoch, "wait_seconds": 0})
    assert r.status_code == 401
    await transport.client.aclose()
    await bad.aclose()
