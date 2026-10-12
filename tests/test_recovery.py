"""Lifecycle and recovery: AT12 (restart), AT13 (cancel/pause), AT24 (desktop workflow, approvals)."""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest

from omca.config import SmtpConfig, config_from_dict
from omca.contracts import TaskRequest, TaskStatus

from .conftest import System, base_config

pytestmark = pytest.mark.browser

CONTACT = {"name": "Asha", "kind": "email", "address": "asha@example.com"}


def office_cfg(tmp_path, site, **extra):
    return config_from_dict(base_config(tmp_path, site, policy={"contacts": [CONTACT]},
                                        limits={"lease_ttl_seconds": 4, "worker_heartbeat_seconds": 1,
                                                "progress_interval_seconds": 3600}, **extra))


def smtp_cfg(sink) -> SmtpConfig:
    return SmtpConfig(host="127.0.0.1", port=sink.port, starttls=False, from_address="owner@example.com")


SEND_PLAN = [
    {"operation": "email.draft", "arguments": {"to": ["asha@example.com"], "subject": "Weekly numbers", "body": "Hi Asha"}},
    {"operation": "email.send", "arguments": {"draft_artifact_id": "$draft", "to": ["asha@example.com"]}},
]


@pytest.mark.acceptance("AT12")
async def test_restart_does_not_repeat_a_completed_send(tmp_path, site, brain, smtp_sink):
    """The worker sent the email but the receipt never reached the coordinator before a crash."""
    cfg = office_cfg(tmp_path, site)
    s1 = await System(cfg, smtp=smtp_cfg(smtp_sink)).start()
    s1.coord.memory.set_preference("tone", "concise")
    brain.office_plan = list(SEND_PLAN)
    original_submit = s1.worker.agent._submit_with_retry

    async def drop_send_receipts(receipt):
        if s1.coord.db.scalar("SELECT operation FROM actions WHERE id=?", (receipt.action_id,)) == "email.send":
            return  # receipt lost in transit
        await original_submit(receipt)

    s1.worker.agent._submit_with_retry = drop_send_receipts
    t = s1.coord.submit_task(TaskRequest(goal="Email Asha the weekly numbers", workflow="desktop_workflow"))
    ap = await s1.pending_approval(t.id)
    s1.coord.decide_approval(ap.approval_id, s1.owner(), "approve")
    for _ in range(100):
        if smtp_sink.messages:
            break
        await asyncio.sleep(0.1)
    assert len(smtp_sink.messages) == 1
    await asyncio.sleep(0.3)
    await s1.stop()  # crash: the coordinator never learned the outcome

    s2 = await System(cfg, smtp=smtp_cfg(smtp_sink)).start()  # same database and worker receipt cache
    t = await s2.wait(t.id, timeout=60)
    assert len(smtp_sink.messages) == 1, "the send was repeated after restart"
    act = s2.coord.db.one("SELECT status FROM actions WHERE task_id=? AND operation='email.send'", (t.id,))
    assert act["status"] == "SUCCEEDED"
    kinds = [e for e in s2.coord.tasks.events(t.id) if e["kind"] == "reconciled"]
    assert kinds and kinds[-1]["via"] == "worker receipt cache"
    # checkpoints persisted: the draft step was not executed again
    assert len([a for a in s2.coord.artifacts.for_task(t.id, "output") if a.name.endswith(".eml")]) == 1
    # stable preferences survive restarts
    assert s2.coord.memory.preferences() == {"tone": "concise"}
    assert t.status == TaskStatus.COMPLETED
    await s2.stop()


@pytest.mark.acceptance("AT12")
async def test_unknown_outcome_is_reconciled_with_the_owner_not_repeated(tmp_path, site, brain, smtp_sink):
    cfg = office_cfg(tmp_path, site)
    s1 = await System(cfg, smtp=smtp_cfg(smtp_sink)).start()
    brain.office_plan = list(SEND_PLAN)
    hang = asyncio.Event()

    async def hanging_send(job, ctx):
        hang.set()
        await asyncio.sleep(3600)

    s1.worker.agent.handlers["email.send"] = hanging_send
    t = s1.coord.submit_task(TaskRequest(goal="Email Asha the weekly numbers", workflow="desktop_workflow"))
    ap = await s1.pending_approval(t.id)
    s1.coord.decide_approval(ap.approval_id, s1.owner(), "approve")
    await asyncio.wait_for(hang.wait(), 30)
    action_id = s1.coord.db.scalar("SELECT id FROM actions WHERE task_id=? AND operation='email.send'", (t.id,))
    await s1.stop()
    # hard crash: the worker never wrote a receipt, only its "started" marker
    worker_dir = Path(cfg.storage.data_dir) / "worker"
    (worker_dir / "receipts" / f"{action_id}.json").unlink(missing_ok=True)
    assert (worker_dir / "started" / f"{action_id}.json").exists()

    s2 = await System(cfg, smtp=smtp_cfg(smtp_sink)).start()
    t2 = await s2.wait(t.id, {TaskStatus.AWAITING_USER}, timeout=60)
    assert "could not confirm" in (t2.blocking_reason or "")
    s2.coord.answer_task(t.id, "not done")
    t3 = await s2.wait(t.id, timeout=60)
    assert smtp_sink.messages == []  # never re-sent automatically
    act = s2.coord.db.one("SELECT status FROM actions WHERE id=?", (action_id,))
    assert act["status"] == "FAILED"
    assert t3.status == TaskStatus.PARTIAL
    await s2.stop()


@pytest.mark.acceptance("AT13")
async def test_pause_stops_new_actions_and_resume_continues(make_system, site, brain):
    s = await make_system()
    brain.browser_script = [{"name": "wait", "arguments": {"ms": 700}} for _ in range(8)] + [
        {"name": "finish", "arguments": {"success": True, "summary": "done"}}]
    brain.verify_markers = ["Contact us"]
    t = s.coord.submit_task(TaskRequest(goal="Wait on the form page", workflow="browser_task",
                                        params={"url": site.url("/form.html")}))
    for _ in range(100):
        if s.coord.tasks.require(t.id).action_count >= 2:
            break
        await asyncio.sleep(0.1)
    s.coord.pause_task(t.id)
    await s.wait(t.id, {TaskStatus.PAUSED}, timeout=10)
    await asyncio.sleep(1.5)
    frozen = s.coord.tasks.require(t.id).action_count
    await asyncio.sleep(1.5)
    assert s.coord.tasks.require(t.id).action_count == frozen, "actions continued while paused"
    assert s.coord.tasks.require(t.id).status == TaskStatus.PAUSED
    s.coord.resume_task(t.id)
    t = await s.wait(t.id, timeout=60)
    assert t.status == TaskStatus.COMPLETED, (str(t.result.get("details", {}).get("verdict")), t.result.get("summary"),
                                              [e["text"] for e in s.coord.tasks.events(t.id) if e["kind"] == "limitation"],
                                              [(e["from"], e["to"]) for e in s.coord.tasks.events(t.id) if e["kind"] == "status"])
    # completed steps were replayed, not repeated
    waits = s.coord.db.scalar("SELECT COUNT(*) FROM actions WHERE task_id=? AND operation='browser.wait'", (t.id,))
    assert waits == 8


@pytest.mark.acceptance("AT13")
async def test_cancel_reports_completed_effects_and_never_completes(make_system, site, brain):
    s = await make_system()
    brain.browser_script = [{"name": "wait", "arguments": {"ms": 600}} for _ in range(10)]
    t = s.coord.submit_task(TaskRequest(goal="Wait on the form page", workflow="browser_task",
                                        params={"url": site.url("/form.html")}))
    for _ in range(100):
        if s.coord.tasks.require(t.id).action_count >= 2:
            break
        await asyncio.sleep(0.1)
    s.coord.cancel_task(t.id)
    t = await s.wait(t.id, timeout=20)
    assert t.status == TaskStatus.CANCELLED
    count = t.action_count
    await asyncio.sleep(1.5)
    assert s.coord.tasks.require(t.id).action_count <= count + 1  # at most the in-flight action finishes
    await s.flush()
    assert any("cancelled" in m and "Completed effects kept" in m for m in s.channel.texts())


@pytest.mark.acceptance("AT13")
async def test_pause_during_approval_wait_then_cancel(make_system, site, brain):
    s = await make_system()
    start = len(site.submissions)
    brain.browser_script = [{"name": "click", "arguments": {"ref_name": "Send message"}}]
    t = s.coord.submit_task(TaskRequest(goal="Use the form", workflow="browser_task", params={"url": site.url("/form.html")}))
    ap = await s.pending_approval(t.id)
    s.coord.pause_task(t.id)
    await s.wait(t.id, {TaskStatus.PAUSED}, timeout=10)
    s.coord.cancel_task(t.id)
    await s.wait(t.id, timeout=10)
    assert s.coord.approvals.get(ap.approval_id).status.value == "INVALIDATED"
    assert len(site.submissions) == start


@pytest.mark.acceptance("AT24")
async def test_desktop_office_workflow_confirms_consequential_changes(make_system, site, brain, smtp_sink, tmp_path):
    from io import BytesIO

    from openpyxl import Workbook, load_workbook

    s = await make_system(smtp=smtp_cfg(smtp_sink), policy={"contacts": [CONTACT]})
    wb = Workbook()
    wb.active["A1"], wb.active["B1"] = "Region", "Total"
    wb.active["A2"], wb.active["B2"] = "North", 10
    buf = BytesIO()
    wb.save(buf)
    t = s.coord.submit_task(TaskRequest(goal="Update totals and email Asha", workflow="desktop_workflow"))
    art = s.coord.artifacts.put(t.id, "input", "sales.xlsx", buf.getvalue()).manifest.artifact_id
    s.coord.tasks.update(t.id, request=t.request.model_copy(update={"attachments": [art]}))
    brain.office_plan = [
        {"operation": "spreadsheet.read", "arguments": {"artifact_id": art}},
        {"operation": "spreadsheet.write", "arguments": {"artifact_id": art, "cells": {"B2": 42, "C1": "Checked"},
                                                        "output_name": "sales-updated.xlsx"}},
        *SEND_PLAN,
    ]
    ap = await s.pending_approval(t.id)
    assert ap.operation == "email.send" and ap.destination == "asha@example.com"
    await asyncio.sleep(0.5)
    assert smtp_sink.messages == []  # nothing sent before the decision
    s.coord.decide_approval(ap.approval_id, s.owner(), "approve")
    t = await s.wait(t.id, timeout=60)
    assert t.status == TaskStatus.COMPLETED, t.result
    assert len(smtp_sink.messages) == 1 and smtp_sink.messages[0]["to"] == ["asha@example.com"]
    outputs = {a.name: a for a in s.coord.artifacts.for_task(t.id, "output")}
    updated = load_workbook(BytesIO(s.coord.artifacts.read_bytes(outputs["sales-updated.xlsx"].artifact_id)))
    assert updated.active["B2"].value == 42 and updated.active["C1"].value == "Checked"
    original = load_workbook(BytesIO(s.coord.artifacts.read_bytes(art)))
    assert original.active["B2"].value == 10  # the original was not overwritten
